#!/usr/bin/env python3
"""
daily_picks.py - the team score betting system, kept up to date every day of the 2026-27 season and emailed to the
owner (asked for Oct 21, 2026; see .github/workflows/picks.yml).

The system: in every game, pick the team whose team score (the page's formula and constants, fitted on all of
2025-26; read straight from dashboard_template.html) was higher in its own previous game.

Each run does whatever is due, so it's safe to run as often as you like (cron-job.org starts it every 15 minutes):
  morning   the first run after 6:00 am Central: score every game that finished since the last run, grade last
            night's picks, then try every rule (a range of lines x a range of differentials x home/away/either) and
            keep the one that would have made the most units this season so far. That rule decides today's picks.
  pre-game  the first run from 2 hours before the day's first game: pick every game of the day at DraftKings' line
            right then (units are graded on that line), rebuild the spreadsheet and email it with today's picks that
            fit the rule, the rule, and every game of the day.
A game that finished without a pick-time line (the pre-game run didn't happen, or no line was posted yet) is
filled in the next morning with its closing line, and the Line source column says so.

    python tools/daily_picks.py run --data picks-data [--now 2026-10-22T21:00:00Z] [--send-now] [--dry-run]

--send-now does the pre-game step and sends the email right away, even if it was already sent today.
--dry-run does everything as if the email went out, but writes it to --data/email.html instead (for testing on a copy of the data).

Everything it keeps lives in --data (the picks-data branch in the workflow, so it never touches the site):
  games.json             every 2026-27 regular-season game that has finished, with both teams' team scores
  picks.json             one row per game from Oct 6, 2026
  rules.json             each day's rule, and whether that day's email went out
  backtest_2025_26.json  last season's backtest (fixed), for its own tab
  team_score_picks.xlsx  the spreadsheet, rebuilt on every change and attached to the email. Tabs: Dynamic Bet Results
                         (only the picks emailed as fitting that day's rule, graded on the emailed line; the owner's
                         real-world record, added at his request), 2026-27 (every game), Daily rule, 2025-26, Notes
Email settings are the post-game email's: SMTP_USER, SMTP_PASSWORD, EMAIL_TO (see postgame_email.py).
"""
from __future__ import annotations

import argparse
import datetime as dt
import html
import json
import os
import re
import smtplib
import sys
import tempfile
import time
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import requests

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import stars_viz as sv  # noqa: E402

TZ = ZoneInfo("America/Chicago")
BETS_FROM = "2026-10-06"   # the first day of bets (the owner's choice); earlier games only count as a previous game
MORNING_HOUR = 6           # Central; last night's games are final by then
LEAD = dt.timedelta(hours=2)
BOOKS = ("DraftKings", "ESPN BET")  # whose line to use, in order
ESPN_ABBR = {"NJ": "NJD", "TB": "TBL", "LA": "LAK", "SJ": "SJS", "UTAH": "UTA"}
ESPN_SB = "https://site.api.espn.com/apis/site/v2/sports/hockey/nhl/scoreboard?dates={:%Y%m%d}"
ESPN_ODDS = "https://sports.core.api.espn.com/v2/sports/hockey/leagues/nhl/events/{0}/competitions/{0}/odds"
XLSX = "team_score_picks.xlsx"
DYNAMIC = "Dynamic Bet Results"


# ---------------------------------------------------------------- the team score (the page's formula, its constants)
def _template_constants() -> tuple[dict, float]:
    t = (ROOT / "dashboard_template.html").read_text(encoding="utf-8")
    team = re.search(r"const TEAM = \{ goal: ([\d.]+), xg: ([\d.]+), cf: ([\d.]+)", t)
    mid = re.search(r"const MID = (\d+)", t)
    return dict(zip(("goal", "xg", "cf"), map(float, team.groups()))), float(mid.group(1))


TEAM, MID = _template_constants()


def team_score(t: dict) -> float:
    pts = TEAM["goal"] * (t["gf"] - t["ga"]) + TEAM["xg"] * (t["xgf"] - t["xga"]) + TEAM["cf"] * (t["cf"] - t["ca"])
    return max(0.0, min(100.0, MID + pts))


def score_game(box: dict, pbp: dict) -> dict:
    """One finished game with both teams' team scores (parsed from the home side, so both teams' shots are in)."""
    t = sv.team_score_inputs(sv.parse_game(box, pbp, {}, box["homeTeam"]["abbrev"]))
    flip = {"gf": t["ga"], "ga": t["gf"], "xgf": t["xga"], "xga": t["xgf"], "cf": t["ca"], "ca": t["cf"]}
    return dict(id=box["id"], date=box["gameDate"], start=box["startTimeUTC"],
                away=box["awayTeam"]["abbrev"], home=box["homeTeam"]["abbrev"],
                away_goals=box["awayTeam"]["score"], home_goals=box["homeTeam"]["score"],
                away_ts=round(team_score(flip), 2), home_ts=round(team_score(t), 2))


# ---------------------------------------------------------------- data in and out
def load(data: Path) -> dict:
    def j(name, default):
        p = data / name
        return json.loads(p.read_text()) if p.exists() else default
    return dict(games=j("games.json", []), picks=j("picks.json", []), rules=j("rules.json", []),
                last_season=j("backtest_2025_26.json", []))


def save(data: Path, st: dict) -> None:
    for name, key in (("games.json", "games"), ("picks.json", "picks"), ("rules.json", "rules")):
        st[key].sort(key=lambda r: (r.get("start") or r.get("date"), r.get("id", 0)))
        (data / name).write_text(json.dumps(st[key], indent=1))


def get_json(url: str):
    for i in range(4):
        try:
            r = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
            if r.status_code == 404:
                return None
            r.raise_for_status()
            return r.json()
        except requests.RequestException:
            if i == 3:
                raise
            time.sleep(2 * (i + 1))


def schedule(day: dt.date) -> list[dict]:
    """The NHL's regular-season games for the week starting on `day`."""
    wk = get_json(f"{sv.WEB}/schedule/{day}") or {}
    return [dict(g, date=w["date"]) for w in wk.get("gameWeek", []) for g in w["games"] if g.get("gameType") == 2]


def american(v) -> int | None:
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    v = str(v).strip().upper()
    return 100 if v in ("EVEN", "EV", "PK") else int(v.replace("+", ""))


def espn_lines(day: dt.date) -> dict:
    """{(away, home): {"away"/"home": {book: {"now": ml, "close": ml}}}} from ESPN's free odds feed."""
    out = {}
    for e in (get_json(ESPN_SB.format(day)) or {}).get("events", []):
        if e.get("season", {}).get("type") != 2:
            continue
        teams = {t["homeAway"]: ESPN_ABBR.get(t["team"]["abbreviation"], t["team"]["abbreviation"])
                 for t in e["competitions"][0]["competitors"]}
        sides = {"away": {}, "home": {}}
        for it in (get_json(ESPN_ODDS.format(e["id"])) or {}).get("items", []):
            book = it.get("provider", {}).get("name", "")
            if book not in BOOKS:
                continue  # also skips ESPN's in-game "Live Odds" entries
            for side in sides:
                o = it.get(f"{side}TeamOdds") or {}
                ml = lambda k: american(((o.get(k) or {}).get("moneyLine") or {}).get("american"))  # noqa: E731
                sides[side][book] = {"now": ml("current") or american(o.get("moneyLine")), "close": ml("close")}
        out[(teams["away"], teams["home"])] = sides
    return out


def pick_line(lines: dict, game: dict, side: str, kind: str) -> tuple[int | None, str | None]:
    books = (lines.get((game["away"], game["home"])) or {}).get(side) or {}
    for book in BOOKS:
        v = (books.get(book) or {}).get(kind)
        if v is not None:
            return v, book
    return None, None


# ---------------------------------------------------------------- picks
def previous(games: list[dict], team: str, before: str) -> float | None:
    """The team's team score in its last finished game that started before `before` (UTC ISO)."""
    best = None
    for g in games:
        if g["start"] < before and team in (g["away"], g["home"]) and (best is None or g["start"] > best["start"]):
            best = g
    return None if best is None else (best["away_ts"] if best["away"] == team else best["home_ts"])


def new_row(game: dict, games: list[dict]) -> dict:
    a, h = previous(games, game["away"], game["start"]), previous(games, game["home"], game["start"])
    pick = None if a is None or h is None or a == h else (game["home"] if h > a else game["away"])
    return dict(id=game["id"], date=game["date"], start=game["start"], away=game["away"], home=game["home"],
                away_prev=a, home_prev=h, pick=pick, line=None, source=None, fits=None, result=None, final=None)


def units(line: int, won: bool) -> float:
    """The owner's rule: a favorite risks enough to win 1, an underdog risks 1."""
    return (line / 100 if line > 0 else 1.0) if won else (line / 100 if line < 0 else -1.0)


def graded(rows: list[dict]) -> list[dict]:
    return [r for r in rows if r["pick"] and r["line"] is not None and r["result"] in ("W", "L")]


# ---------------------------------------------------------------- the rule
def diff_of(r: dict) -> float:
    return round(abs(r["away_prev"] - r["home_prev"]), 2)


def best_rule(rows: list[dict]) -> dict | None:
    """Of every rule "line from lo to hi, differential from dlo to dhi (both included), home picks / away picks /
    either", the one that would have made the most units on these rows (ties: more bets). Every line and every
    differential seen so far is a possible cut point, so nothing better exists within these three conditions.
    Bounds at the edge of what's been seen are left open (None), so a longer line than any so far still counts.

    How: for each range of lines, the differential range is the best run of consecutive differentials (the biggest
    sum of a run, from running totals), worked out for every range of lines at once in numpy."""
    rows = graded(rows)
    if not rows:
        return None
    lg = np.array(sorted({r["line"] for r in rows}))
    dv = np.array(sorted({diff_of(r) for r in rows}))
    li = np.searchsorted(lg, [r["line"] for r in rows])
    di = np.searchsorted(dv, [diff_of(r) for r in rows])
    # units, plus a sliver per bet so that between equal totals the rule with more bets wins (units move in 0.01s)
    key = np.array([units(r["line"], r["result"] == "W") for r in rows]) + 1e-7
    home = np.array([r["pick"] == r["home"] for r in rows])
    a, b = np.triu_indices(len(lg))
    best = (-np.inf,)
    for venue, keep in (("either", np.ones(len(rows), bool)), ("home", home), ("away", ~home)):
        grid = np.zeros((len(lg), len(dv)))
        np.add.at(grid, (li[keep], di[keep]), key[keep])
        by_line = np.vstack([np.zeros(len(dv)), grid.cumsum(0)])
        for s in range(0, len(a), 4000):  # chunks of line ranges, to keep memory small
            aa, bb = a[s:s + 4000], b[s:s + 4000]
            run = np.hstack([np.zeros((len(aa), 1)), (by_line[bb + 1] - by_line[aa]).cumsum(1)])  # running totals
            gain = run[:, 1:] - np.minimum.accumulate(run[:, :-1], axis=1)  # best run ending at each differential
            k = int(np.argmax(gain))
            r_, end = divmod(k, len(dv))
            if gain[r_, end] > best[0] + 1e-9:
                start = int(np.argmin(run[r_, :end + 1]))
                best = (gain[r_, end], int(aa[r_]), int(bb[r_]), start, end, venue)
    _, lo, hi, dlo, dhi, venue = best
    rule = dict(lo=None if lo == 0 else int(lg[lo]), hi=None if hi == len(lg) - 1 else int(lg[hi]),
                dlo=0.0 if dlo == 0 else float(dv[dlo]), dhi=None if dhi == len(dv) - 1 else float(dv[dhi]), venue=venue)
    hits = [r for r in rows if fits(rule, r)]
    rule.update(bets=len(hits), wins=sum(r["result"] == "W" for r in hits),
                units=round(sum(units(r["line"], r["result"] == "W") for r in hits), 2),
                all_bets=len(rows), all_wins=sum(r["result"] == "W" for r in rows),
                all_units=round(sum(units(r["line"], r["result"] == "W") for r in rows), 2))
    return rule


def fits(rule: dict | None, r: dict) -> bool:
    if not rule or not r["pick"] or r["line"] is None:
        return False
    diff = diff_of(r)
    return ((rule["lo"] is None or r["line"] >= rule["lo"]) and (rule["hi"] is None or r["line"] <= rule["hi"])
            and diff >= rule["dlo"] and (rule["dhi"] is None or diff <= rule["dhi"])
            and (rule["venue"] == "either" or (r["pick"] == r["home"]) == (rule["venue"] == "home")))


def fmt_line(v: int) -> str:
    return f"{v:+d}"


def describe(rule: dict | None) -> str:
    if not rule:
        return "No rule yet (no graded picks)."
    lo, hi = rule["lo"], rule["hi"]
    line = ("any line" if lo is None and hi is None else f"any line up to {fmt_line(hi)}" if lo is None
            else f"any line from {fmt_line(lo)} up" if hi is None else f"a line from {fmt_line(lo)} to {fmt_line(hi)}")
    if rule["dlo"] == 0 and rule["dhi"] is None:
        diff = "any differential"
    elif rule["dlo"] == 0:
        diff = f"a differential of {rule['dhi']:g} or less"
    elif rule["dhi"] is None:
        diff = f"a differential of {rule['dlo']:g} or more"
    else:
        diff = f"a differential from {rule['dlo']:g} to {rule['dhi']:g}"
    venue = {"either": "home or away", "home": "home picks only", "away": "away picks only"}[rule["venue"]]
    return f"Bet the pick when it has {line} and {diff}; {venue}."


def record(rule: dict, prefix: str = "") -> str:
    w, n, u = rule[prefix + "wins"], rule[prefix + "bets"], rule[prefix + "units"]
    return f"{w}-{n - w}, {u:+.2f} units"


# ---------------------------------------------------------------- the two daily steps
def score_new_games(st: dict, now: dt.datetime, api: sv.Api) -> int:
    """Adds every regular-season game that has finished since the last one in games.json (looking back at least a week)."""
    have = {g["id"] for g in st["games"]}
    today = now.astimezone(TZ).date()
    last = max((dt.date.fromisoformat(g["date"]) for g in st["games"]), default=today - dt.timedelta(days=7))
    day, added = min(last, today - dt.timedelta(days=7)), 0
    while day <= today:
        for g in schedule(day):
            if g["id"] in have or g.get("gameState") not in sv.FINAL_STATES:
                continue
            box = api.get(sv.url_game(g["id"], "boxscore"), final_only_cache=True)
            pbp = api.get(sv.url_game(g["id"], "play-by-play"), final_only_cache=True)
            if box.get("gameState") in sv.FINAL_STATES:
                st["games"].append(score_game(box, pbp)); have.add(g["id"]); added += 1
        day += dt.timedelta(days=7)
    return added


def morning(st: dict, now: dt.datetime, api: sv.Api) -> dict:
    """Scores and grades everything that finished, then sets today's rule."""
    score_new_games(st, now, api)
    games = sorted(st["games"], key=lambda g: g["start"])
    by_id = {g["id"]: g for g in games}
    rows = {r["id"]: r for r in st["picks"]}
    rules = {r["date"]: r for r in st["rules"]}
    for g in games:  # a game that finished with no row: the pre-game step didn't run that day
        if g["date"] >= BETS_FROM and g["id"] not in rows:
            rows[g["id"]] = new_row(g, games)
            st["picks"].append(rows[g["id"]])
    closing = {}
    for r in st["picks"]:
        g = by_id.get(r["id"])
        if not g or r["result"]:
            continue
        if r["pick"] and r["line"] is None:  # no line at pick time: use the closing line
            day = dt.date.fromisoformat(r["date"])
            for d in (day, day - dt.timedelta(days=1), day + dt.timedelta(days=1)):
                if d not in closing:
                    closing[d] = espn_lines(d)
                line, book = pick_line(closing[d], r, "home" if r["pick"] == r["home"] else "away", "close")
                if line is not None:
                    missed = "no pick-time line" if r["source"] else "missed the pick-time run"
                    r.update(line=line, source=f"{book} closing ({missed})")
                    r["fits"] = fits(rules.get(r["date"]), r) if r["date"] in rules else None
                    break
        home_won = g["home_goals"] > g["away_goals"]
        r["final"] = f'{g["away"]} {g["away_goals"]}, {g["home"]} {g["home_goals"]}'
        if r["pick"]:
            r["result"] = "W" if home_won == (r["pick"] == r["home"]) else "L"
        else:
            r["result"] = "-"
    today = str(now.astimezone(TZ).date())
    rule = best_rule([r for r in st["picks"] if r["date"] < today])
    entry = dict(date=today, rule=rule, emailed=False)
    st["rules"] = [x for x in st["rules"] if x["date"] != today] + [entry]
    return entry


def pregame(st: dict, now: dt.datetime, today_games: list[dict], rule: dict | None) -> list[dict]:
    """Picks every game of the day at the line right now."""
    games = st["games"]
    lines = espn_lines(dt.date.fromisoformat(today_games[0]["date"]))
    rows = {r["id"]: r for r in st["picks"]}
    stamp = now.astimezone(TZ).strftime("%-I:%M %p CT")
    out = []
    for g in sorted(today_games, key=lambda g: (g["startTimeUTC"], g["id"])):
        game = dict(id=g["id"], date=g["date"], start=g["startTimeUTC"], away=g["awayTeam"]["abbrev"], home=g["homeTeam"]["abbrev"])
        r = rows.get(g["id"])
        if r and (r["result"] or r["line"] is not None and "at pick time" in (r["source"] or "")):
            out.append(r); continue  # already picked or graded: a re-send keeps the lines that were emailed
        r2 = new_row(game, games)
        if r2["pick"]:
            line, book = pick_line(lines, game, "home" if r2["pick"] == game["home"] else "away", "now")
            r2.update(line=line, source=f"{book} at pick time ({stamp})" if book else "no line at pick time")
        r2["fits"] = fits(rule, r2)
        if r:
            r.clear(); r.update(r2)
        else:
            st["picks"].append(r2); r = r2
        out.append(r)
    return out


# ---------------------------------------------------------------- the spreadsheet
def workbook(st: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    norm, title = Font(name=F), Font(name=F, bold=True, size=14)
    hdr_fill, hdr_font = PatternFill("solid", fgColor="1F3A2E"), Font(name=F, bold=True, color="FFFFFF")
    green_fill, red_fill = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    thin = Border(bottom=Side(style="thin", color="BBBBBB"))
    big = lambda color=None: Font(name=F, bold=True, size=16, color=color)  # noqa: E731

    def color(ws, ref, size=None):
        f = (lambda c: big(c)) if size else (lambda c: Font(name=F, bold=True, color=c))
        ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"], fill=green_fill, font=f("006100")))
        ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], fill=red_fill, font=f("9C0006")))

    def header(ws, row, heads, widths):
        for c, (t, w) in enumerate(zip(heads, widths), 1):
            x = ws.cell(row, c, t); x.font, x.fill = hdr_font, hdr_fill
            x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = w

    def tiles(ws, items, row=4):
        for i, (lab, f, fmt) in enumerate(items):
            col = 1 + i * 2
            ws.cell(row, col, lab).font = Font(name=F, color="555555", size=10)
            v = ws.cell(row + 1, col, f); v.font, v.number_format = big(), fmt
            v.alignment = Alignment(horizontal="left")

    def season_tab(ws, heading, sub, rows, with_rule):
        H0 = 8
        first, last = H0 + 1, H0 + max(len(rows), 1)
        U, R = f"H{first}:H{last}", f"I{first}:I{last}"
        ws["A1"], ws["A1"].font = heading, title
        ws["A2"] = sub; ws["A2"].font = Font(name=F, italic=True, color="555555")
        items = [("Total units", f"=SUM({U})", "+0.00;-0.00;0.00"),
                 ("Win %", f'=IFERROR(COUNTIF({R},"W")/(COUNTIF({R},"W")+COUNTIF({R},"L")),"")', "0.0%"),
                 ("Record", f'=COUNTIF({R},"W")&"-"&COUNTIF({R},"L")', "@"),
                 ("Bets", f'=COUNTIF({R},"W")+COUNTIF({R},"L")', "0")]
        tiles(ws, items)
        color(ws, "A5", size=True)
        heads = ["Date", "Matchup", "Away team score", "Home team score", "Differential", "Pick", "Line", "Units",
                 "Result", "Final score", "Line source"] + (["Fits that day's rule"] if with_rule else [])
        header(ws, H0, heads, [12, 13, 16, 16, 13, 8, 10, 9, 9, 20, 34, 12][:len(heads)])
        fmts = ["yyyy-mm-dd", "@", "0.0", "0.0", "0.0", "@", "+0;-0;+0", "+0.00;-0.00;0.00", "@", "@", "@", "@"]
        for i, r in enumerate(rows):
            n = first + i
            vals = [dt.date.fromisoformat(r["date"]), f'{r["away"]} @ {r["home"]}', r["away_prev"], r["home_prev"],
                    f'=IF(OR(C{n}="",D{n}=""),"",ABS(C{n}-D{n}))', r["pick"] or "", r["line"],
                    f'=IF(OR(F{n}="",G{n}="",AND(I{n}<>"W",I{n}<>"L")),"",IF(I{n}="W",IF(G{n}>0,G{n}/100,1),IF(G{n}<0,G{n}/100,-1)))',
                    r["result"] if r["result"] in ("W", "L") else "", r["final"] or "", r["source"] or ""]
            if with_rule:
                vals.append({True: "Yes", False: "No", None: ""}[r["fits"]])
            for c, v in enumerate(vals, 1):
                x = ws.cell(n, c, v); x.font, x.number_format, x.border = norm, fmts[c - 1], thin
                x.alignment = Alignment(horizontal="left" if c in (2, 10, 11) else "center")
        color(ws, U)
        ws.freeze_panes = f"A{first}"
        ws.auto_filter.ref = f"A{H0}:{get_column_letter(len(heads))}{last}"

    wb = Workbook(); wb.remove(wb.active)
    cur = sorted(st["picks"], key=lambda r: (r["start"], r["id"]))
    rule_of = {e["date"]: e["rule"] for e in st["rules"]}

    # every pick that was emailed as fitting that day's rule, graded on the line in the email
    dy = wb.create_sheet(DYNAMIC)
    sent = [r for r in cur if r.get("emailed")]
    first, last = 9, 8 + max(len(sent), 1)
    U, R = f"F{first}:F{last}", f"G{first}:G{last}"
    dy["A1"], dy["A1"].font = "Dynamic Bet Results: every pick emailed under that day's rule", title
    dy["A2"] = "Graded on the line in the email. Only picks that were emailed before the game count; nothing is filled in afterwards."
    dy["A2"].font = Font(name=F, italic=True, color="555555")
    tiles(dy, [("Total units", f"=SUM({U})", "+0.00;-0.00;0.00"),
               ("Win %", f'=IFERROR(COUNTIF({R},"W")/(COUNTIF({R},"W")+COUNTIF({R},"L")),"")', "0.0%"),
               ("Record", f'=COUNTIF({R},"W")&"-"&COUNTIF({R},"L")', "@"),
               ("Bets", f'=COUNTIF({R},"W")+COUNTIF({R},"L")', "0"),
               ("Waiting on a result", f'=COUNTA(B{first}:B{last})-COUNTIF({R},"W")-COUNTIF({R},"L")', "0")])
    color(dy, "A5", size=True)
    header(dy, 8, ["Date", "Matchup", "Pick", "Line when emailed", "Differential", "Units", "Result", "Final score",
                   "That day's rule"], [12, 13, 8, 12, 13, 9, 9, 20, 90])
    fmts = ["yyyy-mm-dd", "@", "@", "+0;-0;+0", "0.0", "+0.00;-0.00;0.00", "@", "@", "@"]
    for i, r in enumerate(sent):
        n = first + i
        vals = [dt.date.fromisoformat(r["date"]), f'{r["away"]} @ {r["home"]}', r["pick"], r["line"], diff_of(r),
                f'=IF(AND(G{n}<>"W",G{n}<>"L"),"",IF(G{n}="W",IF(D{n}>0,D{n}/100,1),IF(D{n}<0,D{n}/100,-1)))',
                r["result"] if r["result"] in ("W", "L") else "", r["final"] or "", describe(rule_of.get(r["date"]))]
        for c, v in enumerate(vals, 1):
            x = dy.cell(n, c, v); x.font, x.number_format, x.border = norm, fmts[c - 1], thin
            x.alignment = Alignment(horizontal="left" if c in (2, 8, 9) else "center")
    color(dy, U)
    dy.freeze_panes = f"A{first}"
    dy.auto_filter.ref = f"A8:I{last}"
    ws = wb.create_sheet("2026-27")
    season_tab(ws, "Team score betting system, NHL 2026-27 (from Oct 6, 2026)",
               "Pick the team with the higher team score in its previous game. Updated every morning and before each day's games. See the Notes tab.",
               cur, with_rule=True)
    today_rule = st["rules"][-1]["rule"] if st["rules"] else None
    ws["A3"] = "Today's rule: " + describe(today_rule); ws["A3"].font = Font(name=F, bold=True, color="1F3A2E")

    rt = wb.create_sheet("Daily rule")
    rt["A1"], rt["A1"].font = "Each day's rule: the one that would have made the most units this season up to that morning", title
    header(rt, 3, ["Date", "Rule", "That rule, season to date", "Every pick, season to date", "Picks emailed",
                   "Their record", "Their units"], [12, 70, 22, 22, 10, 12, 11])
    for i, e in enumerate(sorted(st["rules"], key=lambda e: e["date"])):
        n, rule = 4 + i, e["rule"]
        d = dt.date.fromisoformat(e["date"])
        A, H, I_ = (f"'{DYNAMIC}'!$A$9:$A$5000", f"'{DYNAMIC}'!$F$9:$F$5000", f"'{DYNAMIC}'!$G$9:$G$5000")
        vals = [d, describe(rule), record(rule) if rule else "", record(rule, "all_") if rule else "",
                f'=COUNTIFS({A},A{n})',
                f'=COUNTIFS({A},A{n},{I_},"W")&"-"&COUNTIFS({A},A{n},{I_},"L")',
                f'=SUMIFS({H},{A},A{n})']
        for c, v in enumerate(vals, 1):
            x = rt.cell(n, c, v); x.font, x.border = norm, thin
            x.number_format = {1: "yyyy-mm-dd", 7: "+0.00;-0.00;0.00"}.get(c, "@" if c in (2, 3, 4, 6) else "0")
            x.alignment = Alignment(horizontal="left" if c in (2, 3, 4) else "center", wrap_text=c == 2, vertical="top")
    color(rt, "G4:G1000")
    rt.freeze_panes = "A4"

    if st["last_season"]:
        season_tab(wb.create_sheet("2025-26"), "Team score betting backtest, NHL 2025-26 regular season (from Oct 14, 2025)",
                   "Pick the team with the higher team score in its previous game; bet the closing moneyline. See the Notes tab.",
                   st["last_season"], with_rule=False)

    nt = wb.create_sheet("Notes")
    notes = [
        ("How it works", None),
        ("System", "In every game, pick the team whose team score was higher in its own previous game that season."),
        ("Team score", f"The website's formula and constants (dashboard_template.html, teamScore): {MID:g} + {TEAM['goal']} x goal difference (no empty-net goals, no shootout) + {TEAM['xg']} x expected goals difference + {TEAM['cf']} x score-adjusted 5-on-5 shot attempt difference, kept between 0 and 100. The constants were fitted on every game of 2025-26 (all 32 teams) and aren't refitted."),
        ("Units", "Win on a minus line = +1. Win on a plus line = line / 100 (e.g. +240 wins 2.4). Loss on a minus line = line / 100 (e.g. -240 loses 2.4). Loss on a plus line = -1. Results include overtime and the shootout."),
        ("Lines", "From ESPN's free odds feed: DraftKings, or ESPN BET where DraftKings isn't listed. Oct 6-21, 2026 and all of 2025-26 use closing lines. From Oct 22, 2026, every game is picked 2 hours before the day's first game at the line right then, and units are graded on that line. A game without a line at that time gets its closing line, and the Line source column says so."),
        ("Daily schedule", f"Every morning after {MORNING_HOUR} am Central: last night's games are scored and graded, and the rule is reset. Two hours before the day's first game: every game is picked, this file is rebuilt and emailed."),
        ("The daily rule", "Every rule of the form 'line from X to Y, differential from A up to B, home picks / away picks / either' is tried on every graded 2026-27 pick up to that morning, with every line and every differential seen so far as a possible cut point, and the one that would have made the most units wins (ties go to the rule with more bets). A bound at the edge of the lines seen so far is left open. 'Fits that day's rule' marks the picks it chose."),
        ("Caution", "The daily rule is the best fit to the past by construction, so its record always looks good. The honest test is the Dynamic Bet Results tab: only the picks that were emailed under each day's rule, before their games, at the line in the email. On 2025-26, rules found this way didn't carry over from one half of the season to the other or to 2026-27."),
        ("Sources", "Game data: NHL public API (api-web.nhle.com). Odds: ESPN public API. Updated " + dt.datetime.now(TZ).strftime("%Y-%m-%d %H:%M CT") + "."),
    ]
    nt.column_dimensions["A"].width, nt.column_dimensions["B"].width = 18, 110
    for i, (a, b) in enumerate(notes, 1):
        x = nt.cell(i, 1, a); x.font = Font(name=F, bold=True, size=12 if b is None else 10)
        if b is not None:
            y = nt.cell(i, 2, b); y.font = norm; y.alignment = Alignment(wrap_text=True, vertical="top")
            x.alignment = Alignment(vertical="top")
    wb.save(path)


# ---------------------------------------------------------------- the email
def ct(iso: str) -> str:
    return dt.datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(TZ).strftime("%-I:%M %p")


def email(day: dt.date, rows: list[dict], entry: dict, prev: dict | None) -> tuple[str, str, str]:
    rule = entry["rule"]
    hits = [r for r in rows if r["fits"]]
    nice = f"{day:%a %b} {day.day}"
    subj = f"Team score picks, {nice}: {len(hits)} of {len(rows)} game{'s' if len(rows) != 1 else ''} fit the rule"
    num = lambda v, f="{:.1f}": "" if v is None else f.format(v)  # noqa: E731
    diff = lambda r: None if r["away_prev"] is None or r["home_prev"] is None else abs(r["away_prev"] - r["home_prev"])  # noqa: E731
    line = lambda r: "no line yet" if r["pick"] and r["line"] is None else fmt_line(r["line"]) if r["line"] is not None else ""  # noqa: E731
    changed = "" if not prev or not prev.get("rule") else (
        " (same as yesterday)" if describe(prev["rule"]) == describe(rule) else f" (yesterday's was: {describe(prev['rule'])})")

    th = 'style="text-align:left;padding:4px 10px;border-bottom:2px solid #1f3a2e"'
    td = 'style="padding:4px 10px;border-bottom:1px solid #ddd"'

    def table(heads, body):
        return (f'<table style="border-collapse:collapse;font-family:Arial,sans-serif;font-size:14px">'
                f'<tr>{"".join(f"<th {th}>{h}</th>" for h in heads)}</tr>'
                + "".join("<tr>" + "".join(f"<td {td}>{html.escape(str(v))}</td>" for v in row) + "</tr>" for row in body) + "</table>")

    pick_rows = [[ct(r["start"]), f'{r["away"]} @ {r["home"]}', r["pick"], line(r), num(diff(r))] for r in hits]
    all_rows = [[ct(r["start"]), f'{r["away"]} @ {r["home"]}', num(r["away_prev"]), num(r["home_prev"]), num(diff(r)),
                 r["pick"] or "no pick", line(r), "Yes" if r["fits"] else "No"] for r in rows]
    nopick = [r for r in rows if not r["pick"]]
    so_far = (f"This season, that rule would have gone {record(rule)} on {rule['bets']} of {rule['all_bets']} picks "
              f"(every pick: {record(rule, 'all_')}).") if rule else ""
    note = ("Lines are DraftKings' at the time of this email; units are graded on them. All times Central. "
            "The daily rule is the best fit to the season so far, so its past record always looks good: "
            "the spreadsheet's Dynamic Bet Results tab (every pick emailed under the rule, at the emailed line) is the real test.")
    page = f"""<html><body style="font-family:Arial,sans-serif;font-size:15px;color:#222">
<h2 style="color:#1f3a2e">Today's picks that fit the rule</h2>
{table(["Time", "Matchup", "Pick", "Line", "Differential"], pick_rows) if hits else "<p>None of today's games fit the rule.</p>"}
<h2 style="color:#1f3a2e">The rule</h2>
<p><b>{html.escape(describe(rule))}</b>{html.escape(changed)}</p><p>{html.escape(so_far)}</p>
<h2 style="color:#1f3a2e">All of today's games</h2>
{table(["Time", "Matchup", "Away team score", "Home team score", "Differential", "Pick", "Line", "Fits rule"], all_rows)}
{"<p>No pick: " + html.escape(", ".join(f'{r["away"]} @ {r["home"]}' for r in nopick)) + " (a team has no previous game yet, or the two scores are tied).</p>" if nopick else ""}
<p style="font-size:12px;color:#888">Team score = each team's score in its previous game. The full season is in the attached spreadsheet. {html.escape(note)}</p>
</body></html>"""
    text = "\n".join([f"TODAY'S PICKS THAT FIT THE RULE ({nice})"]
                     + ([f"  {a}  {b}: {c} {d} (differential {e})" for a, b, c, d, e in pick_rows] or ["  None of today's games fit the rule."])
                     + ["", "THE RULE", "  " + describe(rule) + changed, "  " + so_far, "", "ALL OF TODAY'S GAMES"]
                     + [f"  {a}  {b}: away {c}, home {d}, differential {e}, pick {f} {g}, fits rule: {h}" for a, b, c, d, e, f, g, h in all_rows]
                     + ["", note])
    return subj, text, page


def send_mail(subj: str, text: str, page: str, attachment: Path) -> str:
    # Google shows app passwords in groups of four ("abcd efgh ijkl mnop"); the spaces aren't part of it
    user, pw = os.environ["SMTP_USER"].strip(), "".join(os.environ["SMTP_PASSWORD"].split())
    to = os.environ.get("EMAIL_TO") or user
    msg = EmailMessage()
    msg["Subject"], msg["From"], msg["To"] = subj, user, to
    msg.set_content(text)
    msg.add_alternative(page, subtype="html")
    msg.add_attachment(attachment.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=attachment.name)
    with smtplib.SMTP_SSL(os.environ.get("SMTP_HOST") or "smtp.gmail.com", int(os.environ.get("SMTP_PORT") or 465), timeout=60) as s:
        s.login(user, pw)
        s.send_message(msg)
    return to


# ---------------------------------------------------------------- one run
def run(data: Path, now: dt.datetime, send_now: bool = False, dry_run: bool = False, api: sv.Api | None = None) -> str:
    data.mkdir(parents=True, exist_ok=True)
    api = api or sv.Api(Path(tempfile.gettempdir()) / "daily_picks_cache")
    st = load(data)
    local = now.astimezone(TZ)
    today = local.date()
    today_games = [g for g in schedule(today) if g["date"] == str(today)]
    starts = [dt.datetime.fromisoformat(g["startTimeUTC"].replace("Z", "+00:00")) for g in today_games]
    first = min(starts, default=None)
    entry = next((e for e in st["rules"] if e["date"] == str(today)), None)
    # due from 2 hours before the first game until the last one starts (after that, the morning fills in closing lines)
    due = bool(today_games) and (send_now or (first - LEAD <= now < max(starts) and not (entry and entry["emailed"])))
    done = []
    if entry is None and (local.hour >= MORNING_HOUR or due):
        entry = morning(st, now, api)
        done.append(f"morning: rule for {today} is: {describe(entry['rule'])}")
    if due:
        rows = pregame(st, now, today_games, entry["rule"])
        workbook(st, data / XLSX)
        prev = next((e for e in sorted(st["rules"], key=lambda e: e["date"], reverse=True) if e["date"] < str(today)), None)
        subj, text, page = email(today, rows, entry, prev)
        if not dry_run and not (os.environ.get("SMTP_USER") and os.environ.get("SMTP_PASSWORD")):
            done.append(f"pre-game: {subj} (not sent: SMTP_USER / SMTP_PASSWORD aren't set)")
        else:
            for r in rows:  # mark first, so the attached sheet's Dynamic Bet Results tab includes today's picks
                r["emailed"] = r.get("emailed") or bool(r["fits"])
            workbook(st, data / XLSX)
            if dry_run:
                (data / "email.html").write_text(page, encoding="utf-8")
                (data / "email.txt").write_text(subj + "\n\n" + text, encoding="utf-8")
                done.append(f"pre-game: {subj} (dry run: written to {data / 'email.html'}, everything else as if sent)")
            else:
                done.append(f"pre-game: {subj} (sent to {send_mail(subj, text, page, data / XLSX)})")
            entry["emailed"] = True
    if done:
        workbook(st, data / XLSX)
        save(data, st)
    return "\n".join(done) or f"nothing due ({local:%Y-%m-%d %H:%M} CT; first game today: {ct(first.isoformat()) + ' CT' if first else 'none'})"


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run", help="do whatever is due")
    r.add_argument("--data", type=Path, required=True)
    r.add_argument("--now", help="pretend it's this time (UTC ISO), for testing")
    r.add_argument("--send-now", action="store_true", help="do the pre-game step and email now, even if already sent")
    r.add_argument("--dry-run", action="store_true", help="as if sent, but write the email to --data/email.html (use on a copy of the data)")
    a = ap.parse_args(argv)
    now = dt.datetime.fromisoformat(a.now.replace("Z", "+00:00")) if a.now else dt.datetime.now(dt.timezone.utc)
    print(run(a.data, now, a.send_now, a.dry_run))


if __name__ == "__main__":
    main()
